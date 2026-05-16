from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = ROOT / "docs"
OUTPUT_PATH = DOCS_DIR / "ASE_Food_Ordering_Final_Report_Template.docx"


ACCENT = RGBColor(25, 79, 135)
LIGHT = RGBColor(235, 242, 250)


def set_page_layout(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_run(run, size=11, bold=False, color=None, font_name="Calibri"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    if color:
        run.font.color.rgb = color


def add_paragraph(doc, text, style=None, align=None, color=None, bold=False, size=11, space_after=6):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    style_run(run, size=16 if level == 1 else 13, bold=True, color=ACCENT)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        style_run(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        style_run(run)


def add_figure_placeholder(doc, title, guidance):
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    table.autofit = True
    top = table.cell(0, 0)
    top.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(top, "EAF2FA")
    p1 = top.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(title)
    style_run(r1, size=11, bold=True, color=ACCENT)

    bottom = table.cell(1, 0)
    bottom.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p2 = bottom.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2 = p2.add_run(guidance)
    style_run(r2, size=10)
    doc.add_paragraph("")


def add_metadata_table(doc):
    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    labels = [
        ("Course", "Advanced Software Engineering - SDEV 4304"),
        ("Project Title", "Food Ordering Web Application using Microservices Architecture"),
        ("Student Name", "Leenah Alborsh"),
        ("Student ID", "220221651"),
        ("Supervisor", "Dr. Abdelkareem Alashqar"),
        ("Semester", "2nd Semester 2025/2026"),
    ]
    for i, (label, value) in enumerate(labels):
        c1, c2 = table.rows[i].cells
        c1.text = ""
        c2.text = ""
        set_cell_shading(c1, "EAF2FA")
        p1 = c1.paragraphs[0]
        p2 = c2.paragraphs[0]
        r1 = p1.add_run(label)
        r2 = p2.add_run(value)
        style_run(r1, bold=True, color=ACCENT)
        style_run(r2)
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def build_doc():
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_page_layout(doc)

    add_paragraph(
        doc,
        "Islamic University of Gaza",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=ACCENT,
        bold=True,
        size=18,
        space_after=4,
    )
    add_paragraph(
        doc,
        "Faculty of Information Technology",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=13,
        bold=True,
        space_after=4,
    )
    add_paragraph(
        doc,
        "Advanced Software Engineering Running Project",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=ACCENT,
        bold=True,
        size=22,
        space_after=10,
    )
    add_paragraph(
        doc,
        "Food Ordering Web Application using Microservices Architecture",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=15,
        space_after=18,
    )
    add_metadata_table(doc)
    add_paragraph(doc, "", space_after=18)
    add_paragraph(
        doc,
        "Submission Date: ____________________",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=12,
        bold=True,
    )

    doc.add_section(WD_SECTION.NEW_PAGE)

    add_paragraph(doc, "Project Links", align=WD_ALIGN_PARAGRAPH.CENTER, color=ACCENT, bold=True, size=20)
    add_paragraph(doc, "Public GitHub Repository Link:", bold=True, size=12)
    add_paragraph(doc, "[Insert the public GitHub repository link here]", size=11)
    add_paragraph(doc, "Public Video Link (Google Drive or YouTube):", bold=True, size=12, space_after=4)
    add_paragraph(doc, "[Insert the public video link here]", size=11, space_after=12)
    add_paragraph(
        doc,
        "Note: Keep both links public and accessible before exporting the final PDF.",
        size=10,
        color=ACCENT,
    )

    doc.add_section(WD_SECTION.NEW_PAGE)

    add_heading(doc, "1. Introduction")
    add_paragraph(
        doc,
        "This project presents a Food Ordering Web Application implemented using a microservices architecture. "
        "The system models the core ordering flow between customers, restaurants, payments, kitchen preparation, "
        "and delivery. The final implementation focuses on a practical, student-friendly solution that can run locally, "
        "demonstrate service communication clearly, and remain aligned with the project design submitted during the semester.",
    )

    add_heading(doc, "2. Problem Description")
    add_paragraph(
        doc,
        "Traditional monolithic systems make it harder to separate responsibilities such as customer management, menu management, "
        "ordering, payment, kitchen processing, and delivery tracking. The Food Ordering System solves this by dividing the domain "
        "into smaller services, where each microservice owns a specific business capability and its related data.",
    )

    add_heading(doc, "3. Project Objectives")
    add_bullets(
        doc,
        [
            "Implement a web-based food ordering system using microservices.",
            "Separate the business domain into clear bounded contexts.",
            "Use Spring Boot and MySQL to build independent services.",
            "Demonstrate communication between services through REST APIs.",
            "Containerize at least one microservice using Docker Desktop.",
            "Prepare a realistic project deliverable that matches the documented design honestly.",
        ],
    )

    add_heading(doc, "4. Bounded Contexts")
    add_numbered(
        doc,
        [
            "Customer Bounded Context: manages customer profiles, addresses, and basic registration data.",
            "Menu Bounded Context: manages menu items, prices, and item availability.",
            "Ordering Bounded Context: manages order creation, quantity, total price, and order lifecycle.",
            "Payment Bounded Context: processes payment requests and stores transaction results.",
            "Kitchen Bounded Context: receives confirmed orders and tracks preparation status.",
            "Delivery Bounded Context: tracks delivery assignment and final delivery status.",
        ],
    )

    add_heading(doc, "5. Final Implemented Architecture")
    add_paragraph(
        doc,
        "The final implementation is organized as a public GitHub monorepo that contains six Spring Boot microservices. "
        "Each service is implemented as an independent Gradle-based application with its own controller, service, repository, entity, "
        "configuration, and database connection settings.",
    )
    add_bullets(
        doc,
        [
            "customer-service",
            "menu-service",
            "ordering-service",
            "payment-service",
            "kitchen-service",
            "delivery-service",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 1. Monorepo Structure",
        "Insert a screenshot from VS Code or File Explorer showing the six service folders inside the project repository.",
    )

    add_heading(doc, "6. Implemented Microservices")
    add_heading(doc, "6.1 Customer Service", level=2)
    add_paragraph(
        doc,
        "Customer Service is responsible for storing and retrieving customer information. It provides REST endpoints for listing customers, "
        "retrieving a customer by ID, and creating new customers. This service is also the service used for the Docker deployment requirement.",
    )
    add_bullets(
        doc,
        [
            "GET /api/customers",
            "GET /api/customers/{id}",
            "POST /api/customers",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 2. Customer Service Endpoint Result",
        "Insert a screenshot from Postman or the browser showing GET http://localhost:9001/api/customers returning data.",
    )

    add_heading(doc, "6.2 Menu Service", level=2)
    add_paragraph(
        doc,
        "Menu Service manages menu items and exposes endpoints to list all menu items, retrieve a single item, and create new items. "
        "Ordering Service uses this service to validate the selected menu item before an order is stored.",
    )
    add_bullets(
        doc,
        [
            "GET /api/menu/items",
            "GET /api/menu/items/{id}",
            "POST /api/menu/items",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 3. Menu Service Endpoint Result",
        "Insert a screenshot from Postman showing GET http://localhost:9003/api/menu/items.",
    )

    add_heading(doc, "6.3 Ordering Service", level=2)
    add_paragraph(
        doc,
        "Ordering Service is the central service in the implemented workflow. It validates the customer and menu item, calculates totals, "
        "sends a payment request, creates a kitchen task after successful payment, and reflects the order status throughout the ordering lifecycle.",
    )
    add_bullets(
        doc,
        [
            "GET /api/orders",
            "GET /api/orders/{id}",
            "POST /api/orders",
            "POST /api/orders/graphql",
            "POST /api/orders/grpc",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 4. Order Creation Response",
        "Insert a screenshot from Postman showing POST http://localhost:9002/api/orders with a successful response that includes order, payment, and kitchen data.",
    )

    add_heading(doc, "6.4 Payment Service", level=2)
    add_paragraph(
        doc,
        "Payment Service simulates payment processing for the educational workflow. It receives payment requests from Ordering Service and stores "
        "the transaction result with the payment method, amount, and payment status.",
    )
    add_bullets(
        doc,
        [
            "GET /api/payments",
            "GET /api/payments/{id}",
            "POST /api/payments/process",
        ],
    )

    add_heading(doc, "6.5 Kitchen Service", level=2)
    add_paragraph(
        doc,
        "Kitchen Service receives confirmed orders after successful payment. It stores a preparation task and allows updating the preparation status, "
        "for example from PENDING_PREPARATION to READY.",
    )
    add_bullets(
        doc,
        [
            "GET /api/kitchen/orders",
            "GET /api/kitchen/orders/{id}",
            "POST /api/kitchen/orders",
            "PUT /api/kitchen/orders/{id}/status",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 5. Kitchen Status Update",
        "Insert a screenshot showing the Kitchen Service status update request where prepStatus becomes READY.",
    )

    add_heading(doc, "6.6 Delivery Service", level=2)
    add_paragraph(
        doc,
        "Delivery Service manages the final stage of the workflow. When the kitchen marks an order as READY, a delivery record can be created and "
        "its status can progress until the order is delivered.",
    )
    add_bullets(
        doc,
        [
            "GET /api/deliveries",
            "GET /api/deliveries/{id}",
            "POST /api/deliveries",
            "PUT /api/deliveries/{id}/status",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 6. Delivery Status Update",
        "Insert a screenshot showing the delivery status update such as OUT_FOR_DELIVERY or DELIVERED.",
    )

    add_heading(doc, "7. Communication Between Services")
    add_paragraph(
        doc,
        "The implemented system mainly uses REST-based synchronous communication. Ordering Service coordinates the main workflow by calling the other services as needed.",
    )
    add_numbered(
        doc,
        [
            "Ordering Service requests customer data from Customer Service.",
            "Ordering Service requests item details from Menu Service.",
            "Ordering Service sends a payment request to Payment Service.",
            "After successful payment, Ordering Service creates a kitchen task in Kitchen Service.",
            "When the kitchen status reaches READY, the delivery process starts in Delivery Service.",
        ],
    )
    add_paragraph(
        doc,
        "In addition to REST, GraphQL and gRPC were preserved between Customer Service and Ordering Service as an extra implementation feature. "
        "However, the main demonstrated workflow in this project relies on REST because it is simpler and more reliable for local testing.",
    )

    add_heading(doc, "8. Database Design")
    add_paragraph(
        doc,
        "Each service is configured with its own MySQL database or schema to support the principle that every microservice owns its data. "
        "This improves separation of concerns and matches the microservices architecture style used in the project.",
    )
    add_bullets(
        doc,
        [
            "customer_service_db",
            "menu_service_db",
            "ordering_service_db",
            "payment_service_db",
            "kitchen_service_db",
            "delivery_service_db",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 7. MySQL Databases for the Services",
        "Insert a screenshot from phpMyAdmin or MySQL Workbench showing the databases or schemas used by the six microservices.",
    )

    add_heading(doc, "9. Main Workflow Scenario")
    add_numbered(
        doc,
        [
            "A customer and a menu item already exist in the system.",
            "The client sends a POST request to Ordering Service to create a new order.",
            "Ordering Service validates the customer through Customer Service.",
            "Ordering Service validates the menu item and retrieves its price through Menu Service.",
            "Ordering Service calculates the total price and calls Payment Service.",
            "If payment succeeds, Ordering Service stores the order and creates a Kitchen Service task.",
            "The kitchen status is updated to READY.",
            "Delivery Service updates the delivery lifecycle until the order becomes DELIVERED.",
        ],
    )
    add_figure_placeholder(
        doc,
        "Figure 8. Final Order Status Result",
        "Insert a screenshot showing GET /api/orders/{id} after the workflow is completed and the order status is updated.",
    )

    add_heading(doc, "10. Running and Testing the Project")
    add_paragraph(
        doc,
        "The project was tested locally using XAMPP MySQL, Spring Boot applications, and Postman. Each service runs on a dedicated port, and the endpoints were tested manually.",
    )
    ports_table = doc.add_table(rows=1, cols=2)
    ports_table.style = "Table Grid"
    hdr = ports_table.rows[0].cells
    hdr[0].text = ""
    hdr[1].text = ""
    set_cell_shading(hdr[0], "EAF2FA")
    set_cell_shading(hdr[1], "EAF2FA")
    style_run(hdr[0].paragraphs[0].add_run("Service"), bold=True, color=ACCENT)
    style_run(hdr[1].paragraphs[0].add_run("Port"), bold=True, color=ACCENT)
    for name, port in [
        ("customer-service", "9001"),
        ("ordering-service", "9002"),
        ("menu-service", "9003"),
        ("payment-service", "9004"),
        ("kitchen-service", "9005"),
        ("delivery-service", "9006"),
    ]:
        row = ports_table.add_row().cells
        row[0].text = name
        row[1].text = port
    doc.add_paragraph("")
    add_paragraph(
        doc,
        "The testing evidence should include screenshots of successful requests from Postman for the customer list, menu list, order creation, kitchen update, delivery update, and final order status.",
    )

    add_heading(doc, "11. Docker Deployment")
    add_paragraph(
        doc,
        "The Docker requirement in this project is satisfied by containerizing at least one microservice, preferably Customer Service. "
        "The Docker build and run steps should be documented honestly based on what is executed successfully on Docker Desktop.",
    )
    add_paragraph(doc, "Docker build command:", bold=True)
    add_paragraph(doc, "docker build -t customer-service .", size=10)
    add_paragraph(doc, "Docker run command:", bold=True)
    add_paragraph(
        doc,
        'docker run --name customer-service-docker -p 9001:9001 -e CUSTOMER_DB_URL="jdbc:mysql://host.docker.internal:3306/customer_service_db?createDatabaseIfNotExist=true&useSSL=false&allowPublicKeyRetrieval=true&serverTimezone=UTC" -e CUSTOMER_DB_USERNAME="root" -e CUSTOMER_DB_PASSWORD="" customer-service',
        size=10,
    )
    add_figure_placeholder(
        doc,
        "Figure 9. Customer Service Running in Docker Desktop",
        "Insert a screenshot from Docker Desktop showing the running customer-service container and its mapped port.",
    )

    add_heading(doc, "12. Implemented Scope")
    add_bullets(
        doc,
        [
            "Six Spring Boot microservices were implemented in one public monorepo.",
            "The system uses MySQL and separates data ownership by service.",
            "The core workflow is implemented through REST communication.",
            "Customer Service and Ordering Service also preserve GraphQL and gRPC as additional communication options.",
            "The project can be tested manually using Postman and browser-based endpoint checks.",
            "Docker support was prepared for at least one microservice.",
        ],
    )

    add_heading(doc, "13. Designed or Future Work")
    add_bullets(
        doc,
        [
            "Kafka-based event-driven communication.",
            "API Gateway for centralized routing.",
            "Consul for service discovery.",
            "Kubernetes for orchestration and scaling.",
            "Full CI/CD pipeline using GitHub Actions and Docker Hub.",
        ],
    )
    add_paragraph(
        doc,
        "These items were part of the broader academic design and can be presented as future improvements unless they are fully implemented and tested later.",
    )

    add_heading(doc, "14. Conclusion")
    add_paragraph(
        doc,
        "This project demonstrates how a Food Ordering System can be divided into multiple independent microservices while keeping the implementation practical and suitable for an academic running project. "
        "The final system applies Spring Boot, MySQL, REST APIs, and Docker preparation to present a realistic microservices solution aligned with the semester documentation.",
    )

    add_heading(doc, "15. Screenshot Checklist Before Final PDF", level=1)
    add_bullets(
        doc,
        [
            "Monorepo structure screenshot.",
            "MySQL databases screenshot.",
            "Customer Service endpoint screenshot.",
            "Menu Service endpoint screenshot.",
            "Order creation response screenshot.",
            "Kitchen READY update screenshot.",
            "Delivery status update screenshot.",
            "Final order status screenshot.",
            "Docker Desktop running container screenshot.",
            "Public GitHub repository screenshot.",
        ],
    )

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_doc()
